from __future__ import annotations

import queue
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

import excel_diff
from excel_block_sorter import (
    DEFAULT_OUTPUT_DIR,
    SortConfig,
    column_to_index,
    columns_to_indices,
    process_files,
    result_line,
)


def default_desktop_output_dir() -> Path:
    desktop = Path.home() / "Desktop"
    if desktop.exists():
        return desktop / "Excel_Sirali_Ciktilar"
    return Path.cwd() / DEFAULT_OUTPUT_DIR


class ExcelBlockSorterApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Excel Blok Sıralama ve Karşılaştırma")
        self.geometry("960x700")
        self.minsize(900, 640)

        self.selected_paths: list[Path] = []
        self.log_queue: queue.Queue[str] = queue.Queue()
        self.worker: threading.Thread | None = None

        # Sıralama Değişkenleri
        self.sort_column_var = tk.StringVar(value="J")
        self.first_row_var = tk.StringVar(value="auto")
        self.sheet_var = tk.StringVar(value="(Tüm Sayfalar)")
        self.output_var = tk.StringVar(value=str(default_desktop_output_dir()))
        self.direction_var = tk.StringVar(value="desc")
        self.suffix_var = tk.StringVar(value="_sirali")
        self.output_name_var = tk.StringVar(value="")
        self.remove_blank_rows_var = tk.BooleanVar(value=False)

        # Karşılaştırma (Diff) Değişkenleri
        self.old_excel_var = tk.StringVar(value="")
        self.new_excel_var = tk.StringVar(value="")
        self.diff_key_col_var = tk.StringVar(value="")
        self.diff_old_sheet_var = tk.StringVar(value="(Tüm Sayfalar)")
        self.diff_new_sheet_var = tk.StringVar(value="(Tüm Sayfalar)")
        self.diff_output_var = tk.StringVar(value=str(default_desktop_output_dir()))

        self._configure_styles()
        self._build_ui()
        self.after(100, self._poll_log_queue)

    def _configure_styles(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass  # Fallback to default if clam is not available

        # Palette: Sleek Blue & Slate Grey
        style.configure(".", background="#f8fafc", foreground="#0f172a", font=("Segoe UI", 10))
        
        # Labels & Headers
        style.configure("TLabel", background="#f8fafc", foreground="#0f172a")
        style.configure("Header.TLabel", background="#f8fafc", foreground="#1e3a8a", font=("Segoe UI", 16, "bold"))
        style.configure("Sub.TLabel", background="#f8fafc", foreground="#475569", font=("Segoe UI", 10))
        
        # Buttons
        style.configure("TButton", font=("Segoe UI", 9, "bold"), padding=6)
        style.configure("Accent.TButton", background="#2563eb", foreground="#ffffff", font=("Segoe UI", 10, "bold"))
        style.map("Accent.TButton",
                  background=[("active", "#1d4ed8"), ("disabled", "#93c5fd")],
                  foreground=[("disabled", "#ffffff")])
        
        # Checkbutton / Radiobutton
        style.configure("TCheckbutton", background="#f8fafc")
        style.configure("TRadiobutton", background="#f8fafc")

        # Notebook
        style.configure("TNotebook", background="#cbd5e1", borderwidth=0)
        style.configure("TNotebook.Tab", background="#e2e8f0", foreground="#475569", padding=(16, 6), font=("Segoe UI", 10))
        style.map("TNotebook.Tab",
                  background=[("selected", "#f8fafc")],
                  foreground=[("selected", "#2563eb")],
                  font=[("selected", ("Segoe UI", 10, "bold"))])

        # LabelFrame
        style.configure("TLabelframe", background="#f8fafc", relief="solid", borderwidth=1, bordercolor="#cbd5e1")
        style.configure("TLabelframe.Label", background="#f8fafc", foreground="#1e3a8a", font=("Segoe UI", 10, "bold"))

    def _build_ui(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)

        # Üst Bilgi Başlığı (Ortak)
        header = ttk.Frame(self, padding=(18, 16, 18, 8))
        header.grid(row=0, column=0, sticky="ew")
        header.columnconfigure(0, weight=1)

        title = ttk.Label(header, text="Excel Blok Sıralama & Karşılaştırma", style="Header.TLabel")
        title.grid(row=0, column=0, sticky="w")
        subtitle = ttk.Label(
            header,
            text="Kurumsal veri tablolarında otomatik blok sıralama yapar ve Excel dosyalarını hücre düzeyinde karşılaştırır.",
            style="Sub.TLabel"
        )
        subtitle.grid(row=1, column=0, sticky="w", pady=(4, 0))

        # Ana Notebook (Sekmeler)
        self.notebook = ttk.Notebook(self)
        self.notebook.grid(row=1, column=0, sticky="nsew", padx=18, pady=8)

        # Sekme 1: Blok Sıralama
        self.tab_sort = ttk.Frame(self.notebook, padding=8)
        self.notebook.add(self.tab_sort, text="Blok Sıralama")
        self._build_sort_tab(self.tab_sort)

        # Sekme 2: Karşılaştırma (Diff)
        self.tab_diff = ttk.Frame(self.notebook, padding=8)
        self.notebook.add(self.tab_diff, text="Excel Karşılaştırma (Diff)")
        self._build_diff_tab(self.tab_diff)

        # Alt Bilgi / Progress Bar (Ortak)
        footer = ttk.Frame(self, padding=(18, 8, 18, 16))
        footer.grid(row=2, column=0, sticky="ew")
        footer.columnconfigure(0, weight=1)

        self.progress = ttk.Progressbar(footer, mode="indeterminate")
        self.progress.grid(row=0, column=0, sticky="ew")

    # -----------------------------------------------------------------------
    # SEKMELERİ İNŞA ETME
    # -----------------------------------------------------------------------

    def _build_sort_tab(self, parent):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(0, weight=1)

        body = ttk.PanedWindow(parent, orient=tk.HORIZONTAL)
        body.grid(row=0, column=0, sticky="nsew")

        left = ttk.Frame(body, padding=10)
        right = ttk.Frame(body, padding=10)
        body.add(left, weight=2)
        body.add(right, weight=3)

        self._build_file_panel(left)
        self._build_settings_panel(right)

    def _build_file_panel(self, parent):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(1, weight=1)

        label = ttk.Label(parent, text="Dosya Seçimi", font=("Segoe UI", 11, "bold"), foreground="#1e3a8a")
        label.grid(row=0, column=0, sticky="w")

        list_frame = ttk.Frame(parent)
        list_frame.grid(row=1, column=0, sticky="nsew", pady=(8, 8))
        list_frame.columnconfigure(0, weight=1)
        list_frame.rowconfigure(0, weight=1)

        self.path_list = tk.Listbox(
            list_frame, height=12, activestyle="none", font=("Segoe UI", 9),
            background="#ffffff", foreground="#0f172a", selectbackground="#3b82f6", selectforeground="#ffffff"
        )
        self.path_list.grid(row=0, column=0, sticky="nsew")
        self.path_list.bind("<<ListboxSelect>>", self._on_path_list_select)

        scrollbar = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.path_list.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        self.path_list.configure(yscrollcommand=scrollbar.set)

        buttons = ttk.Frame(parent)
        buttons.grid(row=2, column=0, sticky="ew")
        buttons.columnconfigure((0, 1, 2), weight=1)

        ttk.Button(buttons, text="Dosya Ekle", command=self.add_files).grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ttk.Button(buttons, text="Klasör Ekle", command=self.add_folder).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(buttons, text="Listeyi Temizle", command=self.clear_paths).grid(row=0, column=2, sticky="ew", padx=(6, 0))

        output_frame = ttk.LabelFrame(parent, text="Çıktı Klasörü", padding=10)
        output_frame.grid(row=3, column=0, sticky="ew", pady=(14, 0))
        output_frame.columnconfigure(0, weight=1)
        
        ttk.Entry(output_frame, textvariable=self.output_var, font=("Segoe UI", 9)).grid(row=0, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(output_frame, text="Masaüstü", command=self.use_desktop_output).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(output_frame, text="Seç", command=self.choose_output).grid(row=0, column=2)

    def _build_settings_panel(self, parent):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(2, weight=1)

        settings = ttk.LabelFrame(parent, text="Sıralama Ayarları", padding=12)
        settings.grid(row=0, column=0, sticky="ew")
        settings.columnconfigure(1, weight=1)
        settings.columnconfigure(3, weight=1)

        ttk.Label(settings, text="Sıralama Kolonu").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=5)
        ttk.Entry(settings, textvariable=self.sort_column_var, width=14).grid(row=0, column=1, sticky="w", pady=5)

        ttk.Label(settings, text="Sayfa Seçimi").grid(row=0, column=2, sticky="w", padx=(18, 8), pady=5)
        self.sheet_combo = ttk.Combobox(settings, textvariable=self.sheet_var, state="readonly")
        self.sheet_combo.grid(row=0, column=3, sticky="ew", pady=5)
        self.sheet_combo['values'] = ["(Tüm Sayfalar)"]
        self.sheet_combo.set("(Tüm Sayfalar)")

        ttk.Label(settings, text="Yön").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=5)
        direction = ttk.Frame(settings)
        direction.grid(row=1, column=1, columnspan=3, sticky="w", pady=5)
        ttk.Radiobutton(direction, text="Büyükten Küçüğe", variable=self.direction_var, value="desc").grid(row=0, column=0)
        ttk.Radiobutton(direction, text="Küçükten Büyüğe", variable=self.direction_var, value="asc").grid(row=0, column=1, padx=(18, 0))

        ttk.Label(settings, text="Yeni Dosya Adı").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=5)
        ttk.Entry(settings, textvariable=self.output_name_var).grid(row=2, column=1, columnspan=3, sticky="ew", pady=5)

        ttk.Label(settings, text="Çoklu Dosya Eki").grid(row=3, column=0, sticky="w", padx=(0, 8), pady=5)
        ttk.Entry(settings, textvariable=self.suffix_var, width=14).grid(row=3, column=1, sticky="w", pady=5)

        ttk.Checkbutton(settings, text="Boş Satırları Sil", variable=self.remove_blank_rows_var).grid(
            row=3, column=2, columnspan=2, sticky="w", padx=(18, 0), pady=5
        )

        help_box = ttk.LabelFrame(parent, text="Çalışma Kuralları", padding=10)
        help_box.grid(row=1, column=0, sticky="ew", pady=(12, 0))
        help_text = (
            "• Başlıkları atlayarak ilk gerçek toplam satırını otomatik tespit eder.\n"
            "• Sıralama kolonları tekli (J) veya çoklu (J, K) girilebilir.\n"
            "• Genel Toplam satırları alt kısımda sabit kalır, ara toplamlar sıralanır.\n"
            "• Satırdaki tüm hücreleri komple taşır.\n"
            "• Detay satırında birleşik hücre bulunursa o blok sıralanmadan güvenle atlanır."
        )
        ttk.Label(help_box, text=help_text, justify="left", font=("Segoe UI", 9)).grid(row=0, column=0, sticky="w")

        log_frame = ttk.LabelFrame(parent, text="Sıralama İşlem Günlüğü", padding=8)
        log_frame.grid(row=2, column=0, sticky="nsew", pady=(12, 0))
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)
        
        self.log_text = tk.Text(
            log_frame, height=8, wrap="word", state="disabled",
            font=("Consolas", 9), background="#ffffff", foreground="#0f172a"
        )
        self.log_text.grid(row=0, column=0, sticky="nsew")
        log_scrollbar = ttk.Scrollbar(log_frame, orient=tk.VERTICAL, command=self.log_text.yview)
        log_scrollbar.grid(row=0, column=1, sticky="ns")
        self.log_text.configure(yscrollcommand=log_scrollbar.set)

        self.run_button = ttk.Button(parent, text="Sıralamayı Başlat", command=self.run_sorting, style="Accent.TButton")
        self.run_button.grid(row=3, column=0, sticky="ew", pady=(10, 0))

    def _build_diff_tab(self, parent):
        parent.columnconfigure(0, weight=1)
        parent.rowconfigure(0, weight=1)

        body = ttk.PanedWindow(parent, orient=tk.HORIZONTAL)
        body.grid(row=0, column=0, sticky="nsew")

        left = ttk.Frame(body, padding=10)
        right = ttk.Frame(body, padding=10)
        body.add(left, weight=2)
        body.add(right, weight=3)

        # Karşılaştırma Sol Panel (Dosyalar)
        left.columnconfigure(0, weight=1)

        files_frame = ttk.LabelFrame(left, text="Karşılaştırılacak Excel Dosyaları", padding=12)
        files_frame.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        files_frame.columnconfigure(1, weight=1)

        ttk.Label(files_frame, text="Eski Excel Dosyası:").grid(row=0, column=0, sticky="w", pady=5, padx=(0, 8))
        ttk.Entry(files_frame, textvariable=self.old_excel_var).grid(row=0, column=1, sticky="ew", pady=5)
        ttk.Button(files_frame, text="Seç", command=self.choose_old_excel, width=6).grid(row=0, column=2, pady=5, padx=(8, 0))

        ttk.Label(files_frame, text="Yeni Excel Dosyası:").grid(row=1, column=0, sticky="w", pady=5, padx=(0, 8))
        ttk.Entry(files_frame, textvariable=self.new_excel_var).grid(row=1, column=1, sticky="ew", pady=5)
        ttk.Button(files_frame, text="Seç", command=self.choose_new_excel, width=6).grid(row=1, column=2, pady=5, padx=(8, 0))

        output_frame = ttk.LabelFrame(left, text="Karşılaştırma Sonuç Klasörü", padding=12)
        output_frame.grid(row=1, column=0, sticky="ew")
        output_frame.columnconfigure(0, weight=1)
        
        ttk.Entry(output_frame, textvariable=self.diff_output_var).grid(row=0, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(output_frame, text="Masaüstü", command=self.use_desktop_diff_output).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(output_frame, text="Seç", command=self.choose_diff_output).grid(row=0, column=2)

        # Karşılaştırma Sağ Panel (Ayarlar & Günlük)
        right.columnconfigure(0, weight=1)
        right.rowconfigure(2, weight=1)

        settings = ttk.LabelFrame(right, text="Analiz ve Karşılaştırma Ayarları", padding=12)
        settings.grid(row=0, column=0, sticky="ew")
        settings.columnconfigure(1, weight=1)

        ttk.Label(settings, text="Anahtar Kolon (Eşleştirme)").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=5)
        ttk.Entry(settings, textvariable=self.diff_key_col_var, width=14).grid(row=0, column=1, sticky="w", pady=5)

        ttk.Label(settings, text="Eski Sayfa").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=5)
        self.diff_old_sheet_combo = ttk.Combobox(settings, textvariable=self.diff_old_sheet_var, state="readonly")
        self.diff_old_sheet_combo.grid(row=1, column=1, sticky="ew", pady=5)
        self.diff_old_sheet_combo['values'] = ["(Tüm Sayfalar)"]
        self.diff_old_sheet_combo.set("(Tüm Sayfalar)")

        ttk.Label(settings, text="Yeni Sayfa").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=5)
        self.diff_new_sheet_combo = ttk.Combobox(settings, textvariable=self.diff_new_sheet_var, state="readonly")
        self.diff_new_sheet_combo.grid(row=2, column=1, sticky="ew", pady=5)
        self.diff_new_sheet_combo['values'] = ["(Tüm Sayfalar)"]
        self.diff_new_sheet_combo.set("(Tüm Sayfalar)")

        guide_frame = ttk.Frame(right, padding=4)
        guide_frame.grid(row=1, column=0, sticky="ew", pady=(6, 0))
        guide_text = (
            "• Anahtar kolon yazılırsa satırlar o kolondaki benzersiz değere göre eşleştirilir.\n"
            "• Boş bırakılırsa karşılaştırma direkt satır numarası sırasıyla yapılır.\n"
            "• Eski ve yeni sayfaları farklı seçerek çapraz sayfa analizi yapabilirsiniz.\n"
            "• Sarı (FFFF00): Hücre değişti. Yeşil (C6EFCE): Yeni satır. Kırmızı (FFC7CE): Silinen satır."
        )
        ttk.Label(guide_frame, text=guide_text, justify="left", font=("Segoe UI", 9), foreground="#475569").grid(row=0, column=0, sticky="w")

        log_frame = ttk.LabelFrame(right, text="Karşılaştırma Sonuç Günlüğü", padding=8)
        log_frame.grid(row=2, column=0, sticky="nsew", pady=(10, 0))
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)
        
        self.diff_log_text = tk.Text(
            log_frame, height=8, wrap="word", state="disabled",
            font=("Consolas", 9), background="#ffffff", foreground="#0f172a"
        )
        self.diff_log_text.grid(row=0, column=0, sticky="nsew")
        log_scrollbar = ttk.Scrollbar(log_frame, orient=tk.VERTICAL, command=self.diff_log_text.yview)
        log_scrollbar.grid(row=0, column=1, sticky="ns")
        self.diff_log_text.configure(yscrollcommand=log_scrollbar.set)

        self.diff_run_button = ttk.Button(right, text="Dosyaları Karşılaştır ve Farkları Renklendir", command=self.run_comparison, style="Accent.TButton")
        self.diff_run_button.grid(row=3, column=0, sticky="ew", pady=(10, 0))

    # -----------------------------------------------------------------------
    # DİNAMİK SAYFA LİSTESİ YÜKLEME (THREADED)
    # -----------------------------------------------------------------------

    def _on_path_list_select(self, event=None):
        selection = self.path_list.curselection()
        if selection:
            idx = selection[0]
            path = self.selected_paths[idx]
            if path.is_file():
                self._update_sheet_combo(path)
                return
        self._update_sheet_combo(None)

    def _update_sheet_combo(self, file_path: Path | None):
        if not file_path or not file_path.is_file() or file_path.suffix.lower() != ".xlsx":
            self.sheet_combo['values'] = ["(Tüm Sayfalar)"]
            self.sheet_combo.set("(Tüm Sayfalar)")
            return

        def load_sheets():
            try:
                wb = load_workbook(file_path, read_only=True, keep_vba=False)
                sheets = wb.sheetnames
                self.after(0, lambda: self._set_sheet_combo_values(sheets))
            except Exception:
                self.after(0, lambda: self._set_sheet_combo_values([]))

        threading.Thread(target=load_sheets, daemon=True).start()

    def _set_sheet_combo_values(self, sheets: list[str]):
        values = ["(Tüm Sayfalar)"] + sheets
        self.sheet_combo['values'] = values
        current = self.sheet_var.get()
        if current not in values:
            self.sheet_combo.set("(Tüm Sayfalar)")

    def _update_diff_old_sheet_combo(self):
        path_str = self.old_excel_var.get().strip()
        if not path_str:
            self.diff_old_sheet_combo['values'] = ["(Tüm Sayfalar)"]
            self.diff_old_sheet_combo.set("(Tüm Sayfalar)")
            return

        file_path = Path(path_str)
        if not file_path.is_file() or file_path.suffix.lower() != ".xlsx":
            self.diff_old_sheet_combo['values'] = ["(Tüm Sayfalar)"]
            self.diff_old_sheet_combo.set("(Tüm Sayfalar)")
            return

        def load_sheets():
            try:
                wb = load_workbook(file_path, read_only=True, keep_vba=False)
                sheets = wb.sheetnames
                self.after(0, lambda: self._set_diff_old_sheet_combo_values(sheets))
            except Exception:
                self.after(0, lambda: self._set_diff_old_sheet_combo_values([]))

        threading.Thread(target=load_sheets, daemon=True).start()

    def _set_diff_old_sheet_combo_values(self, sheets: list[str]):
        values = ["(Tüm Sayfalar)"] + sheets
        self.diff_old_sheet_combo['values'] = values
        current = self.diff_old_sheet_var.get()
        if current not in values:
            self.diff_old_sheet_combo.set("(Tüm Sayfalar)")

    def _update_diff_new_sheet_combo(self):
        path_str = self.new_excel_var.get().strip()
        if not path_str:
            self.diff_new_sheet_combo['values'] = ["(Tüm Sayfalar)"]
            self.diff_new_sheet_combo.set("(Tüm Sayfalar)")
            return

        file_path = Path(path_str)
        if not file_path.is_file() or file_path.suffix.lower() != ".xlsx":
            self.diff_new_sheet_combo['values'] = ["(Tüm Sayfalar)"]
            self.diff_new_sheet_combo.set("(Tüm Sayfalar)")
            return

        def load_sheets():
            try:
                wb = load_workbook(file_path, read_only=True, keep_vba=False)
                sheets = wb.sheetnames
                self.after(0, lambda: self._set_diff_new_sheet_combo_values(sheets))
            except Exception:
                self.after(0, lambda: self._set_diff_new_sheet_combo_values([]))

        threading.Thread(target=load_sheets, daemon=True).start()

    def _set_diff_new_sheet_combo_values(self, sheets: list[str]):
        values = ["(Tüm Sayfalar)"] + sheets
        self.diff_new_sheet_combo['values'] = values
        current = self.diff_new_sheet_var.get()
        if current not in values:
            self.diff_new_sheet_combo.set("(Tüm Sayfalar)")

    # -----------------------------------------------------------------------
    # TETİKLEYİCİ VE YOL SEÇİM FONKSİYONLARI
    # -----------------------------------------------------------------------

    def choose_old_excel(self):
        path = filedialog.askopenfilename(
            title="Eski Excel dosyasını seç",
            filetypes=[("Excel dosyaları", "*.xlsx")]
        )
        if path:
            self.old_excel_var.set(path)
            self._update_diff_old_sheet_combo()

    def choose_new_excel(self):
        path = filedialog.askopenfilename(
            title="Yeni Excel dosyasını seç",
            filetypes=[("Excel dosyaları", "*.xlsx")]
        )
        if path:
            self.new_excel_var.set(path)
            self._update_diff_new_sheet_combo()

    def add_files(self):
        paths = filedialog.askopenfilenames(
            title="Excel dosyası seç",
            filetypes=[("Excel dosyaları", "*.xlsx"), ("Tüm dosyalar", "*.*")]
        )
        self._add_paths(Path(path) for path in paths)

    def add_folder(self):
        folder = filedialog.askdirectory(title="Excel klasörü seç")
        if folder:
            self._add_paths([Path(folder)])

    def clear_paths(self):
        self.selected_paths.clear()
        self.path_list.delete(0, tk.END)
        self._update_sheet_combo(None)

    def choose_output(self):
        folder = filedialog.askdirectory(title="Çıktı klasörünü seç")
        if folder:
            self.output_var.set(folder)

    def choose_diff_output(self):
        folder = filedialog.askdirectory(title="Çıktı klasörünü seç")
        if folder:
            self.diff_output_var.set(folder)

    def use_desktop_output(self):
        self.output_var.set(str(default_desktop_output_dir()))

    def use_desktop_diff_output(self):
        self.diff_output_var.set(str(default_desktop_output_dir()))

    def _add_paths(self, paths):
        existing = {path.resolve() for path in self.selected_paths}
        first_added_file = None
        for path in paths:
            resolved = path.resolve()
            if resolved in existing:
                continue
            self.selected_paths.append(path)
            existing.add(resolved)
            self.path_list.insert(tk.END, str(path))
            if path.is_file() and not first_added_file:
                first_added_file = path

        if first_added_file:
            idx = len(self.selected_paths) - 1
            self.path_list.selection_clear(0, tk.END)
            self.path_list.selection_set(idx)
            self.path_list.activate(idx)
            self._update_sheet_combo(first_added_file)

    # -----------------------------------------------------------------------
    # VALİDASYON VE ÇALIŞTIRMA (SIRALAMA)
    # -----------------------------------------------------------------------

    def validate_config(self) -> tuple[SortConfig, Path, str, str | None] | None:
        if not self.selected_paths:
            messagebox.showerror("Eksik seçim", "En az bir Excel dosyası veya klasörü seçin.")
            return None

        try:
            sort_columns = columns_to_indices(self.sort_column_var.get())
        except ValueError as exc:
            messagebox.showerror("Hatalı kolon", str(exc))
            return None

        output_name = self.output_name_var.get().strip() or None
        if output_name and len(self.selected_paths) != 1:
            messagebox.showerror("Hatalı dosya adı", "Yeni dosya adı sadece tek Excel dosyası seçildiğinde kullanılabilir.")
            return None
        if output_name and self.selected_paths[0].is_dir():
            messagebox.showerror("Hatalı dosya adı", "Klasör seçiminde yeni dosya adı yerine çoklu dosya eki kullanılır.")
            return None

        output_dir = Path(self.output_var.get().strip() or DEFAULT_OUTPUT_DIR)
        suffix = self.suffix_var.get().strip() or "_sirali"
        
        sheet_val = self.sheet_var.get().strip()
        sheet_name = None if sheet_val == "(Tüm Sayfalar)" or not sheet_val else sheet_val

        config = SortConfig(
            sort_column=sort_columns[0],
            first_row=None,
            last_column=None,
            descending=self.direction_var.get() == "desc",
            sheet_name=sheet_name,
            sort_columns=sort_columns,
            remove_blank_rows=self.remove_blank_rows_var.get(),
        )
        return config, output_dir, suffix, output_name

    def run_sorting(self):
        validated = self.validate_config()
        if validated is None:
            return

        config, output_dir, suffix, output_name = validated
        self.run_button.configure(state="disabled")
        self.progress.start(12)
        self._clear_log()
        self._log("İşlem başlatıldı...")

        self.worker = threading.Thread(
            target=self._run_sorting_worker,
            args=(list(self.selected_paths), output_dir, config, suffix, output_name),
            daemon=True,
        )
        self.worker.start()

    def _run_sorting_worker(self, paths, output_dir, config, suffix, output_name):
        try:
            total_results = 0
            skipped_results = 0
            for path in paths:
                self.log_queue.put(f"Taranıyor: {path}")
                selected_output_name = output_name if len(paths) == 1 and path.is_file() else None
                results = process_files(
                    path,
                    output_dir,
                    config,
                    suffix=suffix,
                    output_name=selected_output_name,
                )
                if not results:
                    self.log_queue.put(f"SKIP {path.name}: İşlenecek .xlsx dosyası bulunamadı")
                    skipped_results += 1
                    continue
                for result in results:
                    total_results += 1
                    if result.skipped:
                        skipped_results += 1
                    self.log_queue.put(result_line(result, dry_run=False))
            self.log_queue.put(f"DONE|{total_results}|{skipped_results}")
        except Exception as exc:
            self.log_queue.put(f"ERROR|{exc}")

    # -----------------------------------------------------------------------
    # KARŞILAŞTIRMA (DIFF) ÇALIŞTIRMA
    # -----------------------------------------------------------------------

    def run_comparison(self):
        old_path_str = self.old_excel_var.get().strip()
        new_path_str = self.new_excel_var.get().strip()
        if not old_path_str or not new_path_str:
            messagebox.showerror("Eksik seçim", "Karşılaştırmak için eski ve yeni Excel dosyalarını seçin.")
            return

        old_path = Path(old_path_str)
        new_path = Path(new_path_str)
        if not old_path.exists() or not new_path.exists():
            messagebox.showerror("Hata", "Seçilen Excel dosyaları bulunamadı.")
            return

        key_col = None
        key_col_str = self.diff_key_col_var.get().strip()
        if key_col_str:
            try:
                key_col = column_to_index(key_col_str)
            except ValueError as exc:
                messagebox.showerror("Hatalı kolon", f"Anahtar kolon hatalı: {exc}")
                return

        output_dir = Path(self.diff_output_var.get().strip() or DEFAULT_OUTPUT_DIR)
        
        sheet_old_val = self.diff_old_sheet_var.get().strip()
        sheet_new_val = self.diff_new_sheet_var.get().strip()

        sheet_name_old = None if sheet_old_val == "(Tüm Sayfalar)" or not sheet_old_val else sheet_old_val
        sheet_name_new = None if sheet_new_val == "(Tüm Sayfalar)" or not sheet_new_val else sheet_new_val

        # Validasyon: Biri tümü diğeri spesifik olamaz
        if (sheet_name_old is None and sheet_name_new is not None) or (sheet_name_old is not None and sheet_name_new is None):
            messagebox.showerror(
                "Geçersiz Seçim",
                "İki farklı sayfayı karşılaştırmak için hem Eski hem de Yeni sayfayı tek tek seçmelisiniz.\n"
                "Ya da tüm sayfaları karşılaştırmak için ikisinde de '(Tüm Sayfalar)' seçeneğini bırakmalısınız."
            )
            return

        self.diff_run_button.configure(state="disabled")
        self.progress.start(12)
        self._clear_diff_log()
        self._diff_log("Karşılaştırma işlemi başlatıldı...")

        def worker():
            try:
                self.log_queue.put(f"DIFF_LOG|Karşılaştırılıyor: {old_path.name} <-> {new_path.name}")
                diff_result = excel_diff.diff_workbooks(
                    old_path=old_path,
                    new_path=new_path,
                    output_dir=output_dir,
                    sheet_name_old=sheet_name_old,
                    sheet_name_new=sheet_name_new,
                    key_column=key_col,
                )
                self.log_queue.put(f"DIFF_DONE|{diff_result.output_path}")
                # Detayları logla
                for res in diff_result.sheet_results:
                    if sheet_name_old and sheet_name_new and sheet_name_old != sheet_name_new:
                        label = f"{sheet_name_old} (Eski) vs {sheet_name_new} (Yeni)"
                    else:
                        label = res.sheet_name
                    self.log_queue.put(
                        f"DIFF_LOG|Sayfa: {label} -> "
                        f"{res.changed_cells} değişen hücre, "
                        f"{res.added_rows} eklenen satır, {res.removed_rows} silinen satır"
                    )
                if diff_result.added_sheets:
                    self.log_queue.put(f"DIFF_LOG|Eklenen sayfalar: {', '.join(diff_result.added_sheets)}")
                if diff_result.removed_sheets:
                    self.log_queue.put(f"DIFF_LOG|Silinen sayfalar: {', '.join(diff_result.removed_sheets)}")
            except Exception as exc:
                self.log_queue.put(f"DIFF_ERROR|{exc}")

        threading.Thread(target=worker, daemon=True).start()

    # -----------------------------------------------------------------------
    # LOG QUEUE POLLING & UI UPDATE
    # -----------------------------------------------------------------------

    def _poll_log_queue(self):
        while True:
            try:
                message = self.log_queue.get_nowait()
            except queue.Empty:
                break

            if message.startswith("DONE|"):
                _, total, skipped = message.split("|", 2)
                self._finish_run(int(total), int(skipped))
            elif message.startswith("ERROR|"):
                self._finish_run(0, 1, message.split("|", 1)[1])
            elif message.startswith("DIFF_DONE|"):
                output_path = message.split("|", 1)[1]
                self._finish_diff_run(output_path)
            elif message.startswith("DIFF_ERROR|"):
                self._finish_diff_run(None, error=message.split("|", 1)[1])
            elif message.startswith("DIFF_LOG|"):
                self._diff_log(message.split("|", 1)[1])
            else:
                self._log(message)

        self.after(100, self._poll_log_queue)

    def _finish_run(self, total, skipped, error=None):
        self.progress.stop()
        self.run_button.configure(state="normal")
        if error:
            self._log(f"HATA: {error}")
            messagebox.showerror("İşlem hatası", error)
            return

        self._log(f"İşlem bitti. Toplam: {total}, atlanan/sorunlu: {skipped}.")
        if skipped:
            messagebox.showwarning("İşlem tamamlandı", f"Sıralama bitti. {skipped} dosya/blok atlandı (Detaylar için günlüğe bakın).")
        else:
            messagebox.showinfo("İşlem tamamlandı", "Tüm dosyalar başarıyla sıralandı.")

    def _finish_diff_run(self, output_path, error=None):
        self.progress.stop()
        self.diff_run_button.configure(state="normal")
        if error:
            self._diff_log(f"HATA: {error}")
            messagebox.showerror("Karşılaştırma hatası", error)
            return

        self._diff_log(f"\nKarşılaştırma tamamlandı. Fark raporu: {output_path}")
        messagebox.showinfo("Karşılaştırma tamamlandı", f"İki dosya başarıyla karşılaştırıldı ve renklendirildi.\n\nFark Raporu: {output_path}")

    def _clear_log(self):
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", tk.END)
        self.log_text.configure(state="disabled")

    def _log(self, message):
        self.log_text.configure(state="normal")
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)
        self.log_text.configure(state="disabled")

    def _clear_diff_log(self):
        self.diff_log_text.configure(state="normal")
        self.diff_log_text.delete("1.0", tk.END)
        self.diff_log_text.configure(state="disabled")

    def _diff_log(self, message):
        self.diff_log_text.configure(state="normal")
        self.diff_log_text.insert(tk.END, message + "\n")
        self.diff_log_text.see(tk.END)
        self.diff_log_text.configure(state="disabled")


def main():
    app = ExcelBlockSorterApp()
    app.mainloop()


if __name__ == "__main__":
    main()
