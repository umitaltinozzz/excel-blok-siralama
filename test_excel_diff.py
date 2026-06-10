from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from excel_diff import (
    CellDiff,
    DiffResult,
    FILL_ADDED,
    FILL_CHANGED,
    FILL_REMOVED,
    SheetDiffResult,
    diff_sheets,
    diff_workbooks,
    write_diff_workbook,
)


# ---------------------------------------------------------------------------
# Yardımcı fonksiyonlar
# ---------------------------------------------------------------------------

def _create_workbook(path: Path, data: dict[str, list[list]]) -> Path:
    """Verilen sayfa adı → satır listesi eşlemesinden bir Excel dosyası oluşturur."""
    wb = Workbook()
    first = True
    for sheet_name, rows in data.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


def _fill_color(cell) -> str | None:
    """Hücrenin arka plan rengini döndürür (solid dolgu için)."""
    if cell.fill and cell.fill.fill_type == "solid":
        color = cell.fill.fgColor
        if color and color.rgb and isinstance(color.rgb, str):
            # openpyxl bazen 'AARRGGBB' formatında döndürür
            rgb = color.rgb
            if len(rgb) == 8:
                return rgb[2:]
            return rgb
    return None


# ---------------------------------------------------------------------------
# Testler
# ---------------------------------------------------------------------------

def test_identical_workbooks_produce_no_diffs(tmp_path: Path):
    """Aynı içerikli iki dosya karşılaştırılırsa sıfır fark bulunmalı."""
    data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25], ["Veli", 30]]}
    old = _create_workbook(tmp_path / "eski.xlsx", data)
    new = _create_workbook(tmp_path / "yeni.xlsx", data)

    result = diff_workbooks(old, new)

    assert len(result.sheet_results) == 1
    sheet = result.sheet_results[0]
    assert sheet.changed_cells == 0
    assert sheet.added_rows == 0
    assert sheet.removed_rows == 0
    assert len(sheet.diffs) == 0


def test_changed_cells_are_detected(tmp_path: Path):
    """Üç hücre değiştirilirse tam olarak üç fark tespit edilmeli."""
    old_data = {"Sayfa1": [["Ad", "Yas", "Sehir"], ["Ali", 25, "Ankara"], ["Veli", 30, "Istanbul"]]}
    new_data = {"Sayfa1": [["Ad", "Yas", "Sehir"], ["Ali", 26, "Izmir"], ["Veli", 31, "Istanbul"]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)

    result = diff_workbooks(old, new)

    sheet = result.sheet_results[0]
    assert sheet.changed_cells == 3
    assert sheet.added_rows == 0
    assert sheet.removed_rows == 0

    changed_diffs = [d for d in sheet.diffs if d.diff_type == "changed"]
    assert len(changed_diffs) == 3

    # Kontrol: Ali'nin yaşı 25 → 26
    age_diff = next(d for d in changed_diffs if d.old_value == 25)
    assert age_diff.new_value == 26


def test_added_rows_are_detected(tmp_path: Path):
    """Yeni dosyada fazladan satır varsa eklenen satır olarak tespit edilmeli."""
    old_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25]]}
    new_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25], ["Veli", 30], ["Ayse", 28]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)

    result = diff_workbooks(old, new)

    sheet = result.sheet_results[0]
    assert sheet.added_rows == 2
    assert sheet.changed_cells == 0
    assert sheet.removed_rows == 0

    added_diffs = [d for d in sheet.diffs if d.diff_type == "added"]
    assert len(added_diffs) > 0


def test_removed_rows_are_detected(tmp_path: Path):
    """Eski dosyada olup yeni dosyada olmayan satırlar silinen satır olarak tespit edilmeli."""
    old_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25], ["Veli", 30], ["Ayse", 28]]}
    new_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)

    result = diff_workbooks(old, new)

    sheet = result.sheet_results[0]
    assert sheet.removed_rows == 2
    assert sheet.changed_cells == 0
    assert sheet.added_rows == 0


def test_diff_output_has_colored_cells(tmp_path: Path):
    """Çıktı dosyasında değişen hücreler sarı, eklenen satırlar yeşil olmalı."""
    old_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25]]}
    new_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 30], ["Veli", 28]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)
    output_dir = tmp_path / "cikti"

    result = diff_workbooks(old, new, output_dir)

    assert result.output_path is not None
    assert result.output_path.exists()

    wb = load_workbook(result.output_path)
    ws = wb["Sayfa1"]

    # Ali'nin yaşı değişti → sarı (FFFF00)
    changed_cell = ws.cell(2, 2)  # Satır 2, Kolon B (Yas)
    assert _fill_color(changed_cell) == "FFFF00"
    assert changed_cell.comment is not None
    assert "Eski" in changed_cell.comment.text

    # Veli eklendi → yeşil (C6EFCE)
    added_cell = ws.cell(3, 1)
    assert _fill_color(added_cell) == "C6EFCE"


def test_key_column_matching(tmp_path: Path):
    """Anahtar kolon eşleştirmesiyle satır sırası değişse bile doğru eşleşme yapılmalı."""
    # Eski: Ali satır 2, Veli satır 3
    old_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25], ["Veli", 30]]}
    # Yeni: Sıra değişti – Veli satır 2, Ali satır 3, Ayse satır 4 (eklenen)
    new_data = {"Sayfa1": [["Ad", "Yas"], ["Veli", 35], ["Ali", 25], ["Ayse", 28]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)

    # Anahtar kolon = 1 (Ad kolonu)
    result = diff_workbooks(old, new, key_column=1)

    sheet = result.sheet_results[0]

    # Veli'nin yaşı 30 → 35 olarak değişti
    changed = [d for d in sheet.diffs if d.diff_type == "changed"]
    assert len(changed) == 1
    assert changed[0].old_value == 30
    assert changed[0].new_value == 35

    # Ayse eklendi
    assert sheet.added_rows == 1

    # Ali değişmedi (satır sırası farklı ama anahtar eşleşmesi doğru)
    ali_diffs = [d for d in sheet.diffs if d.diff_type == "changed" and d.old_value == 25]
    assert len(ali_diffs) == 0


def test_added_sheet_detected(tmp_path: Path):
    """Yeni dosyada olup eski dosyada olmayan sayfalar tespit edilmeli."""
    old_data = {"Sayfa1": [["Veri", 1]]}
    new_data = {"Sayfa1": [["Veri", 1]], "YeniSayfa": [["Yeni", 2]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)

    result = diff_workbooks(old, new)

    assert "YeniSayfa" in result.added_sheets
    assert len(result.removed_sheets) == 0


def test_removed_sheet_detected(tmp_path: Path):
    """Eski dosyada olup yeni dosyada olmayan sayfalar tespit edilmeli."""
    old_data = {"Sayfa1": [["Veri", 1]], "EskiSayfa": [["Eski", 2]]}
    new_data = {"Sayfa1": [["Veri", 1]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)

    result = diff_workbooks(old, new)

    assert "EskiSayfa" in result.removed_sheets
    assert len(result.added_sheets) == 0


def test_summary_sheet_created(tmp_path: Path):
    """Çıktı dosyasında 'Fark Özeti' sayfası bulunmalı ve doğru veriler içermeli."""
    old_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 25]]}
    new_data = {"Sayfa1": [["Ad", "Yas"], ["Ali", 30], ["Veli", 28]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)
    output_dir = tmp_path / "cikti"

    result = diff_workbooks(old, new, output_dir)

    wb = load_workbook(result.output_path)
    assert "Fark Özeti" in wb.sheetnames

    ws = wb["Fark Özeti"]
    # Başlık kontrolü
    assert ws.cell(1, 1).value == "Sayfa Adı"
    assert ws.cell(1, 2).value == "Değişen Hücre/Formül"
    assert ws.cell(1, 3).value == "Eklenen Satır"
    assert ws.cell(1, 4).value == "Silinen Satır"

    # İlk sayfa sonuçları
    assert ws.cell(2, 1).value == "Sayfa1"
    assert ws.cell(2, 2).value == 1   # 1 değişen hücre (yaş)
    assert ws.cell(2, 3).value == 1   # 1 eklenen satır (Veli)
    assert ws.cell(2, 4).value == 0   # 0 silinen satır


def test_cross_sheet_comparison(tmp_path: Path):
    """Farklı isimlerdeki iki sayfanın karşılaştırılabildiğini doğrular."""
    old_data = {"SayfaEski": [["Ad", "Yas"], ["Ali", 25]]}
    new_data = {"SayfaYeni": [["Ad", "Yas"], ["Ali", 30]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)
    output_dir = tmp_path / "cikti"

    result = diff_workbooks(
        old,
        new,
        output_dir=output_dir,
        sheet_name_old="SayfaEski",
        sheet_name_new="SayfaYeni"
    )

    assert len(result.sheet_results) == 1
    sheet = result.sheet_results[0]
    assert sheet.sheet_name == "SayfaYeni"
    assert sheet.changed_cells == 1
    assert len(sheet.diffs) == 1
    assert sheet.diffs[0].old_value == 25
    assert sheet.diffs[0].new_value == 30


def test_formula_changes_detected(tmp_path: Path):
    """Formül değişikliklerinin ve formül-değer dönüşümlerinin tespit edildiğini doğrular."""
    old_data = {"Sayfa1": [["Ad", "Toplam", "X", "Y"], ["Ali", "=SUM(C2:D2)", 10, 20], ["Veli", 100, 30, 40]]}
    new_data = {"Sayfa1": [["Ad", "Toplam", "X", "Y"], ["Ali", "=C2+D2", 10, 20], ["Veli", "=SUM(C3:D3)", 30, 40]]}

    old = _create_workbook(tmp_path / "eski.xlsx", old_data)
    new = _create_workbook(tmp_path / "yeni.xlsx", new_data)

    result = diff_workbooks(old, new)
    sheet = result.sheet_results[0]

    assert sheet.changed_cells == 2

    ali_diff = next(d for d in sheet.diffs if d.row == 2 and d.col == 2)
    assert "Formül Değişti" in ali_diff.comment
    assert "=SUM(C2:D2)" in ali_diff.comment
    assert "=C2+D2" in ali_diff.comment

    veli_diff = next(d for d in sheet.diffs if d.row == 3 and d.col == 2)
    assert "Formül Eklendi" in veli_diff.comment
    assert "=SUM(C3:D3)" in veli_diff.comment


def test_format_changes_detected(tmp_path: Path):
    """Sayı biçimlendirmesi (Format) değişikliklerinin tespit edildiğini doğrular."""
    old_data = {"Sayfa1": [["Ad", "Tutar"], ["Ali", 1000]]}
    old = _create_workbook(tmp_path / "eski.xlsx", old_data)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sayfa1"
    ws.append(["Ad", "Tutar"])
    ws.append(["Ali", 1000])
    ws.cell(2, 2).number_format = "$#,##0.00"
    new = tmp_path / "yeni.xlsx"
    wb.save(new)

    result = diff_workbooks(old, new)
    sheet = result.sheet_results[0]

    assert sheet.changed_cells == 1
    diff = sheet.diffs[0]
    assert "Biçimlendirme Değişti" in diff.comment


