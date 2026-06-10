from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side

from excel_block_sorter import (
    SortConfig,
    column_to_index,
    columns_to_indices,
    detect_last_column,
    parse_number,
    process_files,
    sort_sheet,
)


def mark_total_row(ws, row, sort_col):
    cell = ws.cell(row, sort_col)
    cell.fill = PatternFill("solid", fgColor="BFBFBF")
    cell.border = Border(top=Side(style="double"), bottom=Side(style="double"))


def create_sample_workbook(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"

    ws.append(["Grup", "Ad", "J", "K"])
    ws.append(["G1", "TOPLAM", 100, 7])
    ws.append(["G1", "Dusuk", 20, 300])
    ws.append(["G1", "Yuksek", 80, 100])
    ws.append(["G1", "Bos", None, 200])
    ws.append(["G2", "TOPLAM", 50, 9])
    ws.append(["G2", "Orta", 30, 1])
    ws.append(["G2", "Buyuk", 40, 2])

    mark_total_row(ws, 2, 3)
    mark_total_row(ws, 6, 3)
    wb.save(path)


def test_parse_number_supports_excel_and_turkish_text_formats():
    assert parse_number(367349745.58) == 367349745.58
    assert parse_number("367.349.746") == 367349746
    assert parse_number("264.863.878,02 TL") == 264863878.02
    assert parse_number("117 milyon") == 117_000_000
    assert parse_number("1,5 milyon") == 1_500_000
    assert parse_number("-2.700") == -2700
    assert parse_number("2026\nFiili") is None
    assert parse_number("metin") is None


def test_column_to_index_accepts_letters_and_numbers():
    assert column_to_index("J") == 10
    assert column_to_index("n") == 14
    assert column_to_index("3") == 3


def test_columns_to_indices_accepts_multiple_columns():
    assert columns_to_indices("J,K") == (10, 11)
    assert columns_to_indices("J; K 12") == (10, 11, 12)


def test_detect_last_column_keeps_at_least_sort_column(tmp_path):
    source = tmp_path / "sample.xlsx"
    create_sample_workbook(source)

    ws = load_workbook(source)["Rapor"]

    assert detect_last_column(ws, 3) == 4
    assert detect_last_column(ws, 8) == 8
    assert detect_last_column(ws, (3, 8)) == 8


def test_sort_sheet_keeps_total_rows_and_sorts_details_by_selected_column(tmp_path):
    source = tmp_path / "sample.xlsx"
    create_sample_workbook(source)

    wb = load_workbook(source)
    ws = wb["Rapor"]
    changed_blocks, total_rows, detail_rows, skipped_ranges = sort_sheet(
        ws,
        SortConfig(
            sort_column=3,
            first_row=2,
            last_column=None,
            descending=True,
        ),
    )

    assert changed_blocks == 2
    assert total_rows == 2
    assert detail_rows == 5
    assert ws.cell(2, 2).value == "TOPLAM"
    assert ws.cell(3, 2).value == "Yuksek"
    assert ws.cell(4, 2).value == "Dusuk"
    assert ws.cell(5, 2).value == "Bos"
    assert ws.cell(6, 2).value == "TOPLAM"
    assert ws.cell(7, 2).value == "Buyuk"
    assert ws.cell(8, 2).value == "Orta"


def test_sort_column_setting_changes_order(tmp_path):
    source = tmp_path / "sample.xlsx"
    create_sample_workbook(source)

    wb = load_workbook(source)
    ws = wb["Rapor"]
    ws.cell(3, 5).value = "dusuk-ek"
    ws.cell(4, 5).value = "yuksek-ek"
    sort_sheet(
        ws,
        SortConfig(
            sort_column=4,
            first_row=2,
            last_column=None,
            descending=True,
            sort_total_blocks=False,
        ),
    )

    assert [ws.cell(row, 2).value for row in range(3, 6)] == ["Dusuk", "Bos", "Yuksek"]
    assert ws.cell(5, 5).value == "yuksek-ek"


def test_multiple_sort_columns_use_secondary_column_for_ties(tmp_path):
    source = tmp_path / "sample.xlsx"
    create_sample_workbook(source)

    wb = load_workbook(source)
    ws = wb["Rapor"]
    ws.cell(3, 3).value = 20
    ws.cell(4, 3).value = 20
    ws.cell(3, 4).value = 100
    ws.cell(4, 4).value = 300

    sort_sheet(
        ws,
        SortConfig(
            sort_column=3,
            first_row=2,
            descending=True,
            sort_columns=(3, 4),
        ),
    )

    assert [ws.cell(row, 2).value for row in range(3, 5)] == ["Yuksek", "Dusuk"]


def test_total_blocks_are_sorted_by_total_value_and_grand_total_stays_last(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"
    ws.append(["Grup", "Ad", "J", "K"])
    ws.append(["G1", "TOPLAM KUCUK", 50, 0])
    ws.append(["G1", "Kucuk Detay", 50, 0])
    ws.append(["G2", "TOPLAM BUYUK", 100, 0])
    ws.append(["G2", "Buyuk Detay", 100, 0])
    ws.append(["", "Genel Toplam", 150, 0])
    mark_total_row(ws, 2, 3)
    mark_total_row(ws, 4, 3)
    mark_total_row(ws, 6, 3)

    sort_sheet(
        ws,
        SortConfig(
            sort_column=3,
            first_row=None,
            descending=True,
        ),
    )

    assert [ws.cell(row, 2).value for row in range(2, 7)] == [
        "TOPLAM BUYUK",
        "Buyuk Detay",
        "TOPLAM KUCUK",
        "Kucuk Detay",
        "Genel Toplam",
    ]


def test_auto_first_row_ignores_styled_report_header(tmp_path):
    source = tmp_path / "sample.xlsx"
    create_sample_workbook(source)

    wb = load_workbook(source)
    ws = wb["Rapor"]
    ws.insert_rows(1)
    ws.cell(1, 3).value = "2026\nFiili"
    mark_total_row(ws, 1, 3)

    changed_blocks, total_rows, detail_rows, skipped_ranges = sort_sheet(
        ws,
        SortConfig(
            sort_column=3,
            first_row=None,
            last_column=None,
            descending=True,
        ),
    )

    assert changed_blocks == 2
    assert total_rows == 2
    assert detail_rows == 5
    assert ws.cell(1, 3).value == "2026\nFiili"
    assert ws.cell(4, 2).value == "Yuksek"


def test_process_files_writes_copy_without_changing_source(tmp_path):
    source = tmp_path / "sample.xlsx"
    output_dir = tmp_path / "out"
    create_sample_workbook(source)

    results = process_files(
        tmp_path,
        output_dir,
        SortConfig(sort_column=3, first_row=None),
    )

    assert len(results) == 1
    assert results[0].output_path == output_dir / "sample_sirali.xlsx"
    assert results[0].output_path.exists()

    original = load_workbook(source)["Rapor"]
    sorted_ws = load_workbook(results[0].output_path)["Rapor"]

    assert original.cell(3, 2).value == "Dusuk"
    assert sorted_ws.cell(3, 2).value == "Yuksek"


def test_process_files_can_use_output_name_for_single_file(tmp_path):
    source = tmp_path / "sample.xlsx"
    output_dir = tmp_path / "out"
    create_sample_workbook(source)

    results = process_files(
        source,
        output_dir,
        SortConfig(sort_column=3, first_row=None),
        output_name="rapor_sirali",
    )

    assert len(results) == 1
    assert results[0].output_path == output_dir / "rapor_sirali.xlsx"
    assert results[0].output_path.exists()


def test_gui_module_can_be_imported():
    import excel_block_sorter_gui

    assert hasattr(excel_block_sorter_gui, "ExcelBlockSorterApp")


def test_remove_blank_rows_deletes_empty_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"

    ws.append(["Grup", "Ad", "J", "K"])
    ws.append(["G1", "TOPLAM", 100, 7])
    ws.append(["G1", "Dusuk", 20, 300])
    ws.append(["", "", "", ""])  # Empty row
    ws.append(["G1", "Yuksek", 80, 100])
    ws.append(["", "", "", ""])  # Empty row

    mark_total_row(ws, 2, 3)

    changed_blocks, total_rows, detail_rows, skipped_ranges = sort_sheet(
        ws,
        SortConfig(
            sort_column=3,
            first_row=2,
            descending=True,
            remove_blank_rows=True,
        )
    )

    assert ws.max_row == 4
    assert ws.cell(3, 2).value == "Yuksek"
    assert ws.cell(4, 2).value == "Dusuk"


def test_merged_cell_block_is_skipped(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"

    # Block 1 (has merged cell in details)
    ws.append(["G1", "TOPLAM", 100, 7]) # row 1
    ws.append(["G1", "A", 20, 300])      # row 2
    ws.append(["G1", "B", 80, 100])      # row 3
    # Block 2 (clean)
    ws.append(["G2", "TOPLAM", 50, 9])  # row 4
    ws.append(["G2", "C", 10, 1])        # row 5
    ws.append(["G2", "D", 40, 2])        # row 6

    mark_total_row(ws, 1, 3)
    mark_total_row(ws, 4, 3)

    # Merge cells in row 2 (columns 1 and 2)
    ws.merge_cells("A2:B2")

    changed_blocks, total_rows, detail_rows, skipped_ranges = sort_sheet(
        ws,
        SortConfig(
            sort_column=3,
            first_row=1,
            descending=True,
        )
    )

    assert any("birlesik hucre iceriyor, bu blok atlandi" in text for text in skipped_ranges)
    # Block 1 should not be sorted: A remains at row 2, B remains at row 3
    assert ws.cell(2, 1).value == "G1"
    assert ws.cell(3, 2).value == "B"
    # Block 2 should be sorted: D (40) is sorted before C (10)
    assert ws.cell(5, 2).value == "D"
    assert ws.cell(6, 2).value == "C"


def test_total_blocks_with_merged_cells_are_kept_fixed_but_others_are_sorted(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"

    # 3 Toplam Bloğu
    ws.append(["G1", "TOPLAM 1", 50, 0])  # row 1 (temiz)
    ws.append(["G1", "Detay 1", 50, 0])

    ws.append(["G2", "TOPLAM 2", 100, 0]) # row 3 (birleşik hücre içerecek)
    ws.append(["G2", "Detay 2", 100, 0])

    ws.append(["G3", "TOPLAM 3", 200, 0]) # row 5 (temiz)
    ws.append(["G3", "Detay 3", 200, 0])

    mark_total_row(ws, 1, 3)
    mark_total_row(ws, 3, 3)
    mark_total_row(ws, 5, 3)

    # G2 bloğunda (row 3 ve 4) birleşik hücre yapalım
    ws.merge_cells("A3:B3")

    changed_blocks, total_rows, detail_rows, skipped_ranges = sort_sheet(
        ws,
        SortConfig(
            sort_column=3,
            first_row=1,
            descending=True,
        )
    )

    # G2 bloğu yerinde sabit kalmalı çünkü birleşik hücre içeriyor.
    # Ancak G1 ve G3 kendi aralarında yer değiştirebilmeli (200 > 50 olduğu için G3 en başa gelmeli).
    # Normal sıralama sırası şöyle olmalı:
    # 1. G3 bloğu (temiz, en yüksek değere sahip) -> Satır 1-2'ye taşınacak
    # 2. G2 bloğu (sabit kalmalı, yani satır 3-4'te kalacak)
    # 3. G1 bloğu (temiz, en düşük değere sahip) -> Satır 5-6'ya taşınacak

    assert any("birlesik hucre iceriyor, bu toplam blogunun sirasi sabit tutuldu" in text for text in skipped_ranges)

    # Satır 1: G3 TOPLAM 3
    assert ws.cell(1, 2).value == "TOPLAM 3"
    # Satır 3: G2 TOPLAM 2 (sabit kaldı) - A3 hücresi değer tutar, B3 MergedCell'dir
    assert ws.cell(3, 1).value == "G2"
    assert ws.cell(3, 3).value == 100
    # Satır 5: G1 TOPLAM 1
    assert ws.cell(5, 2).value == "TOPLAM 1"


